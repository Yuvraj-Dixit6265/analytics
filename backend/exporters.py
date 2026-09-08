"""exporters.py — turn one report's computed results into a PDF, a PPTX, or
an XLSX. Two very different kinds of PDF live here:

  - build_browser_pdf(url)   — a pixel-accurate capture of the dashboard as
    it actually renders, taken with headless Chromium. Same numbers, same
    boxes, same colours as the live page, because it IS the live page.
    Needs a real URL to visit (the published link), so it only works for
    reports that are actually published.

  - build_table_pdf(name, sections) — a plain, print-friendly listing built
    from computed values, no browser involved. Used for the separate
    "Table" export, and as the editor's own PDF until that also gets a
    browser-backed version.

PPTX and XLSX both build from the same normalized `sections` structure (see
app.py:_render_boxes), so a box that's wrong in one format is wrong — and
fixable — in exactly one place.
"""

import io
import os
import re


# --------------------------------------------------------------------------
# BROWSER PDF — a screenshot of the real page, with clickable section links
# --------------------------------------------------------------------------
def build_browser_pdf(url):
    """Render the published page in headless Chromium, capture it, split it
    into A4 pages, and add invisible clickable links over the on-page
    section-navigation buttons so each one jumps to its section in the PDF.
    The visual result is the dashboard itself — nothing is redrawn."""
    from PIL import Image
    from playwright.sync_api import sync_playwright
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import (
        ArrayObject, DictionaryObject, FloatObject, NameObject, NullObject,
        NumberObject,
    )

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-setuid-sandbox"],
        )
        page = browser.new_page(
            viewport={"width": 1440, "height": 900},
            device_scale_factor=3,
        )

        response = page.goto(url, wait_until="networkidle")
        print("PDF BROWSER URL:", url)
        print("PDF PAGE STATUS:", response.status if response else None)
        print("PDF PAGE URL:", page.url)

        page.wait_for_selector(".sheet", timeout=30000)
        page.wait_for_timeout(1500)
        page.evaluate("""
            async () => { if (document.fonts) { await document.fonts.ready; } }
        """)

        # Hide the download buttons themselves — they shouldn't appear
        # inside the exported file.
        page.add_style_tag(content=".export-actions { display: none !important; }")

        sheet = page.locator(".sheet")
        dimensions = sheet.evaluate("""
            el => {
                const r = el.getBoundingClientRect();
                return { x: r.x, y: r.y, width: r.width, height: r.height };
            }
        """)
        print("PDF SHEET CSS DIMENSIONS:", dimensions)

        # ------------------------------------------------------------
        # Section navigation button positions — become invisible
        # clickable areas in the final PDF.
        # ------------------------------------------------------------
        section_links = sheet.evaluate("""
            sheet => {
                const buttons = Array.from(sheet.querySelectorAll(".section-nav button"));
                const sections = Array.from(
                    sheet.querySelectorAll('.sec[id^="section-"]')
                ).filter(el => el.offsetParent !== null);
                const sr = sheet.getBoundingClientRect();
                return buttons.map((button, index) => {
                    const target = sections[index];
                    if (!target) return null;
                    const br = button.getBoundingClientRect();
                    const tr = target.getBoundingClientRect();
                    return {
                        text: (button.textContent || "").trim() || "Untitled Section",
                        button: {
                            x: br.left - sr.left, y: br.top - sr.top,
                            width: br.width, height: br.height,
                        },
                        targetY: tr.top - sr.top,
                    };
                }).filter(Boolean);
            }
        """)
        print("PDF SECTION LINKS:", section_links)
        print("PDF SECTION LINK COUNT:", len(section_links))
        if not section_links:
            print("NOTE: no section-nav buttons found — the PDF will still "
                  "render, just without in-PDF jump links.")
            print("BUTTON COUNT:", sheet.locator(".section-nav button").count())
            print("SECTION COUNT:", sheet.locator('.sec[id^="section-"]').count())

        screenshot = sheet.screenshot(type="png", animations="disabled")
        device_scale_factor = page.evaluate("() => window.devicePixelRatio")
        print("DEVICE SCALE FACTOR:", device_scale_factor)
        browser.close()

    # ------------------------------------------------------------------
    # Convert the screenshot into A4-sized pages
    # ------------------------------------------------------------------
    image = Image.open(io.BytesIO(screenshot)).convert("RGB")
    print("PDF SCREENSHOT PIXELS:", image.size)

    A4_WIDTH = 2480
    A4_HEIGHT = 3508
    scale = A4_WIDTH / image.width
    scaled_width = A4_WIDTH
    scaled_height = round(image.height * scale)
    image = image.resize((scaled_width, scaled_height), Image.Resampling.LANCZOS)
    print("PDF SCALED IMAGE:", image.size)

    pages = []
    for top in range(0, scaled_height, A4_HEIGHT):
        bottom = min(top + A4_HEIGHT, scaled_height)
        page_image = Image.new("RGB", (A4_WIDTH, A4_HEIGHT), "white")
        crop = image.crop((0, top, scaled_width, bottom))
        page_image.paste(crop, (0, 0))
        pages.append(page_image)

    output = io.BytesIO()
    pages[0].save(output, format="PDF", resolution=300.0, save_all=True,
                  append_images=pages[1:])
    pdf_bytes = output.getvalue()

    # ------------------------------------------------------------------
    # Add invisible internal PDF links over each section-nav button
    # ------------------------------------------------------------------
    reader = PdfReader(io.BytesIO(pdf_bytes))
    writer = PdfWriter()
    for pdf_page in reader.pages:
        writer.add_page(pdf_page)

    A4_WIDTH_PT = 595.28
    A4_HEIGHT_PT = 841.89
    FINAL_WIDTH_PX = pages[0].width
    FINAL_HEIGHT_PX = pages[0].height
    print("FINAL IMAGE DIMENSIONS:", FINAL_WIDTH_PX, "x", FINAL_HEIGHT_PX)

    PX_TO_PT_X = A4_WIDTH_PT / FINAL_WIDTH_PX
    PX_TO_PT_Y = A4_HEIGHT_PT / FINAL_HEIGHT_PX

    css_to_screenshot = float(device_scale_factor)
    css_to_final = scale * css_to_screenshot
    print("CSS TO FINAL:", css_to_final)

    print("\n=== SECTION LINKS DEBUG ===")
    for i, link in enumerate(section_links):
        print(i + 1, "| button:", link.get("button"), "| targetY:", link.get("targetY"))
    print("===========================\n")

    for link in section_links:
        button = link["button"]
        target_y_css = link["targetY"]

        button_x1_px = button["x"] * css_to_final
        button_x2_px = (button["x"] + button["width"]) * css_to_final
        button_top_px = button["y"] * css_to_final
        button_bottom_px = (button["y"] + button["height"]) * css_to_final
        target_y_px = target_y_css * css_to_final

        target_page_index = int(target_y_px // FINAL_HEIGHT_PX)
        target_page_index = max(0, min(target_page_index, len(writer.pages) - 1))

        button_page_index = int(button_top_px // FINAL_HEIGHT_PX)
        if button_page_index < 0 or button_page_index >= len(writer.pages):
            continue

        button_page_top_px = button_page_index * FINAL_HEIGHT_PX
        button_top_page_px = button_top_px - button_page_top_px
        button_bottom_page_px = button_bottom_px - button_page_top_px

        button_x1_pt = button_x1_px * PX_TO_PT_X
        button_x2_pt = button_x2_px * PX_TO_PT_X
        button_y1_pt = A4_HEIGHT_PT - (button_bottom_page_px * PX_TO_PT_Y)
        button_y2_pt = A4_HEIGHT_PT - (button_top_page_px * PX_TO_PT_Y)

        button_x1_pt = max(0.0, min(A4_WIDTH_PT, button_x1_pt))
        button_x2_pt = max(0.0, min(A4_WIDTH_PT, button_x2_pt))
        button_y1_pt = max(0.0, min(A4_HEIGHT_PT, button_y1_pt))
        button_y2_pt = max(0.0, min(A4_HEIGHT_PT, button_y2_pt))

        target_page_top_px = target_page_index * FINAL_HEIGHT_PX
        target_y_page_px = target_y_px - target_page_top_px
        target_y_pt = A4_HEIGHT_PT - (target_y_page_px * PX_TO_PT_Y)
        target_y_pt = max(0.0, min(A4_HEIGHT_PT, target_y_pt))

        print(
            "Creating link:",
            "| button page =", button_page_index + 1,
            "| target page =", target_page_index + 1,
            "| rect =", (round(button_x1_pt, 2), round(button_y1_pt, 2),
                         round(button_x2_pt, 2), round(button_y2_pt, 2)),
            "| targetY =", round(target_y_pt, 2),
        )

        annotation = DictionaryObject()
        annotation.update({
            NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Subtype"): NameObject("/Link"),
            NameObject("/Rect"): ArrayObject([
                FloatObject(button_x1_pt), FloatObject(button_y1_pt),
                FloatObject(button_x2_pt), FloatObject(button_y2_pt),
            ]),
            NameObject("/Border"): ArrayObject([
                NumberObject(0), NumberObject(0), NumberObject(0),
            ]),
            NameObject("/A"): DictionaryObject({
                NameObject("/S"): NameObject("/GoTo"),
                NameObject("/D"): ArrayObject([
                    writer.pages[target_page_index].indirect_reference,
                    NameObject("/XYZ"), NullObject(),
                    FloatObject(target_y_pt), NullObject(),
                ]),
            }),
        })

        annotation_ref = writer._add_object(annotation)
        pdf_page_obj = writer.pages[button_page_index]
        if "/Annots" not in pdf_page_obj:
            pdf_page_obj[NameObject("/Annots")] = ArrayObject()
        pdf_page_obj[NameObject("/Annots")].append(annotation_ref)

    print("\n=== VERIFYING PDF ANNOTATIONS ===")
    total_annotations = 0
    for i, pg in enumerate(writer.pages):
        annots = pg.get("/Annots")
        count = len(annots) if annots else 0
        total_annotations += count
        print("PAGE", i + 1, "ANNOTATIONS:", count)
    print("TOTAL PDF ANNOTATIONS:", total_annotations)

    final_output = io.BytesIO()
    writer.write(final_output)
    final_pdf = final_output.getvalue()
    print("\nFINAL PDF SIZE:", len(final_pdf))
    return final_pdf


# --------------------------------------------------------------------------
# shared small helpers
# --------------------------------------------------------------------------
def _esc(s):
    return (str(s) if s is not None else "").replace("&", "&amp;") \
        .replace("<", "&lt;").replace(">", "&gt;")


def _cell(v):
    return "—" if v is None else str(v)


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------
def build_pptx(name, sections):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(2.9), Inches(12), Inches(1.5))
    p = tb.text_frame.paragraphs[0]
    p.text = name
    p.font.size = Pt(40)
    p.font.bold = True

    for sec in sections:
        slide = prs.slides.add_slide(blank)
        title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8))
        tp = title.text_frame.paragraphs[0]
        tp.text = sec["name"]
        tp.font.size = Pt(28)
        tp.font.bold = True

        y = 1.3
        for box in sec["boxes"]:
            if y > 6.6:
                slide = prs.slides.add_slide(blank)
                y = 0.5

            if box.get("error"):
                tb = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(12), Inches(0.6))
                tb.text_frame.text = f"{box['title']}: {box['error']}"
                y += 0.7
                continue

            if box["kind"] == "value":
                card = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(y), Inches(5.8), Inches(1.5))
                card.fill.solid()
                card.fill.fore_color.rgb = RGBColor(255, 255, 255)
                card.line.color.rgb = RGBColor(213, 221, 218)

                tf = card.text_frame
                tf.clear()
                tf.word_wrap = True

                p1 = tf.paragraphs[0]
                p1.text = box["title"] or ""
                p1.font.size = Pt(12)
                p1.font.color.rgb = RGBColor(0x5A, 0x66, 0x63)

                p2 = tf.add_paragraph()
                p2.text = box.get("text", "—")
                p2.font.size = Pt(30)
                p2.font.bold = True

                if box.get("note"):
                    p3 = tf.add_paragraph()
                    p3.text = box["note"]
                    p3.font.size = Pt(8)
                    p3.font.color.rgb = RGBColor(120, 120, 120)

                y += 1.7

            elif box["kind"] in ("chart", "table") and box.get("rows"):
                rows = box["rows"][:13]
                r, c = len(rows), max(len(x) for x in rows)
                label = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(6), Inches(0.4))
                label.text_frame.text = box["title"] or ""
                label.text_frame.paragraphs[0].font.bold = True
                y += 0.45
                height = min(0.32 * r, 4.8)
                shape = slide.shapes.add_table(r, c, Inches(0.5), Inches(y),
                                                Inches(12.3), Inches(height))
                tbl = shape.table
                for ri in range(r):
                    for ci in range(c):
                        v = rows[ri][ci] if ci < len(rows[ri]) else None
                        cell = tbl.cell(ri, ci)
                        cell.text = "—" if v is None else str(v)
                        cell.text_frame.paragraphs[0].font.size = Pt(11)
                y += height + 0.35

            elif box.get("text"):
                nb = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(12.3), Inches(1))
                nb.text_frame.word_wrap = True
                nb.text_frame.text = f"{box['title']}: {box['text']}" if box["title"] else box["text"]
                y += 1

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------
def build_xlsx(name, sections):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    used = set()

    for sec in sections:
        base = re.sub(r"[\[\]\*\?/\\:]", "", sec["name"] or "Section")[:28] or "Section"
        sheet_name, i = base, 2
        while sheet_name in used:
            sheet_name = f"{base[:26]}-{i}"
            i += 1
        used.add(sheet_name)
        ws = wb.create_sheet(sheet_name)

        row = 1
        for box in sec["boxes"]:
            ws.cell(row=row, column=1, value=box["title"] or "Untitled").font = Font(bold=True)
            row += 1
            if box.get("error"):
                ws.cell(row=row, column=1, value=f"Error: {box['error']}")
                row += 2
                continue
            if box["kind"] == "value":
                ws.cell(row=row, column=1, value=box.get("text", "—"))
                if box.get("note"):
                    ws.cell(row=row, column=2, value=box["note"])
                row += 2
            elif box["kind"] in ("chart", "table") and box.get("rows"):
                for ri, r in enumerate(box["rows"]):
                    for ci, v in enumerate(r):
                        cell = ws.cell(row=row + ri, column=1 + ci, value=v)
                        if ri == 0:
                            cell.font = Font(bold=True)
                            cell.fill = PatternFill("solid", fgColor="E7ECEA")
                row += len(box["rows"]) + 2
            elif box.get("text"):
                ws.cell(row=row, column=1, value=box["text"])
                row += 2
        for col_idx in range(1, 9):
            ws.column_dimensions[chr(64 + col_idx)].width = 22

    if not wb.sheetnames:
        wb.create_sheet("Report")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------
def build_csv(name, sections):
    import csv

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([name])

    for sec in sections:
        writer.writerow([])
        writer.writerow([sec["name"]])
        if sec.get("desc"):
            writer.writerow([sec["desc"]])

        for box in sec["boxes"]:
            writer.writerow([])
            writer.writerow([box["title"] or "Untitled"])
            if box.get("error"):
                writer.writerow(["Error", box["error"]])
            elif box["kind"] == "value":
                writer.writerow(["Value", box.get("text", "—")])
                if box.get("note"):
                    writer.writerow(["Note", box["note"]])
            elif box["kind"] in ("chart", "table") and box.get("rows"):
                for row in box["rows"]:
                    writer.writerow(["—" if v is None else v for v in row])
            elif box.get("text"):
                writer.writerow([box["text"]])

    return buf.getvalue().encode("utf-8")


# --------------------------------------------------------------------------
# TABLE PDF — a plain, print-friendly listing (no browser involved)
# --------------------------------------------------------------------------
def build_table_pdf(name, sections):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=36, bottomMargin=36,
                             leftMargin=36, rightMargin=36)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("TableReportTitle", parent=styles["Title"],
                                  fontSize=20, spaceAfter=20)
    section_style = ParagraphStyle("TableSection", parent=styles["Heading2"],
                                    fontSize=14, spaceBefore=14, spaceAfter=8)
    header_style = ParagraphStyle("TableHeader", parent=styles["BodyText"],
                                   fontSize=9, fontName="Helvetica-Bold", textColor=colors.white)
    cell_style = ParagraphStyle("TableCell", parent=styles["BodyText"], fontSize=9, leading=12)

    story = [Paragraph(_esc(name), title_style)]

    for sec in sections:
        story.append(Paragraph(_esc(sec.get("name", "Untitled Section")), section_style))
        if sec.get("desc"):
            story.append(Paragraph(_esc(sec["desc"]), cell_style))
            story.append(Spacer(1, 8))

        table_data = [[
            Paragraph("Metric", header_style),
            Paragraph("Value", header_style),
            Paragraph("Note / Details", header_style),
        ]]

        for box in sec.get("boxes", []):
            if box.get("error"):
                table_data.append([
                    Paragraph(_esc(box.get("title") or "Error"), cell_style),
                    Paragraph("Error", cell_style),
                    Paragraph(_esc(box["error"]), cell_style),
                ])
            elif box.get("kind") == "value":
                table_data.append([
                    Paragraph(_esc(box.get("title") or "Untitled"), cell_style),
                    Paragraph(_esc(box.get("text", "—")), cell_style),
                    Paragraph(_esc(box.get("note", "")), cell_style),
                ])
            elif box.get("kind") in ("chart", "table") and box.get("rows"):
                rows = box["rows"]
                table_data.append([
                    Paragraph(f"<b>{_esc(box.get('title') or 'Table')}</b>", cell_style), "", "",
                ])
                for row in rows:
                    row_values = ["—" if v is None else str(v) for v in row]
                    if len(row_values) == 1:
                        table_data.append([Paragraph(_esc(row_values[0]), cell_style), "", ""])
                    elif len(row_values) == 2:
                        table_data.append([
                            Paragraph(_esc(row_values[0]), cell_style),
                            Paragraph(_esc(row_values[1]), cell_style), "",
                        ])
                    else:
                        table_data.append([
                            Paragraph(_esc(row_values[0]), cell_style),
                            Paragraph(_esc(row_values[1]), cell_style),
                            Paragraph(_esc(" | ".join(row_values[2:])), cell_style),
                        ])
            elif box.get("text"):
                table_data.append([
                    Paragraph(_esc(box.get("title") or "Details"), cell_style),
                    Paragraph(_esc(box.get("text", "")), cell_style), "",
                ])

        if len(table_data) > 1:
            table = Table(table_data, colWidths=[3.2 * inch, 2.0 * inch, 4.0 * inch],
                           repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D6E62")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D5DDDA")),
                ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("ALIGN", (1, 1), (1, -1), "CENTER"),
            ]))
            story.append(table)
            story.append(Spacer(1, 16))

    doc.build(story)
    return buf.getvalue()